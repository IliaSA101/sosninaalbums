# -*- coding: utf-8 -*-
"""Низкоуровневые помощники docx: XML-примитивы, нумерация списков, инлайн-разметка."""
import re

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor

import mdparse

# Фирменный цвет выносок (не цвет программы) — тот же, что в annotate.py
ACCENT = "D81F26"
ACCENT_RGB = RGBColor(0xD8, 0x1F, 0x26)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x66, 0x66, 0x66)
ZEBRA = "FBF3F3"
CODEBG = "F5F5F5"
BORDER = "D9C6C7"
FONT = "Calibri"
MONO = "Consolas"


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(r'<w:shd %s w:val="clear" w:fill="%s"/>' % (nsdecls("w"), fill)))


def cell_borders(cell):
    cell._tc.get_or_add_tcPr().append(parse_xml(
        '<w:tcBorders %s>' % nsdecls("w")
        + ''.join('<w:%s w:val="single" w:sz="4" w:color="%s"/>' % (s, BORDER)
                  for s in ("top", "bottom", "left", "right"))
        + '</w:tcBorders>'))


def bottom_border(paragraph, color=BORDER, size=6):
    paragraph._p.get_or_add_pPr().append(parse_xml(
        '<w:pBdr %s><w:bottom w:val="single" w:sz="%d" w:space="4" w:color="%s"/></w:pBdr>'
        % (nsdecls("w"), size, color)))


def no_proof(run):
    """Без проверки правописания: коды, пути и адреса не подчёркиваются волной."""
    run._r.get_or_add_rPr().append(parse_xml("<w:noProof %s/>" % nsdecls("w")))


# Порядок дочерних элементов по схеме OOXML: XML-вставки в конец ломают его, а Word
# на нарушении порядка может сообщить о повреждённом файле.
_ORDER = {
    "pPr": ("pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr "
            "suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap "
            "overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid "
            "spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
            "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange"),
    "rPr": ("rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss "
            "imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs "
            "highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout "
            "specVanish oMath"),
    "tblPr": ("tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize "
              "tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook"),
    "trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader",
}


def normalize_order(root):
    """Приводит pPr, rPr, tblPr, trPr поддерева к порядку схемы (стабильно)."""
    for name, order in _ORDER.items():
        idx = {n: i for i, n in enumerate(order.split())}
        for el in root.iter(qn("w:" + name)):
            kids = list(el)
            kids.sort(key=lambda k: idx.get(k.tag.rsplit("}", 1)[-1], len(idx)))
            for k in kids:
                el.remove(k)
            for k in kids:
                el.append(k)


def add_hyperlink(paragraph, text, url):
    """Кликабельная внешняя ссылка."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    rPr.append(parse_xml("<w:noProof %s/>" % nsdecls("w")))
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_field(paragraph, instr, placeholder=""):
    """Word-поле (PAGE / NUMPAGES / TOC …)."""
    def run_with(child):
        r = OxmlElement("w:r")
        r.append(child)
        return r

    def fld(kind):
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        return run_with(fc)

    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    parts = [fld("begin"), run_with(it), fld("separate")]
    if placeholder:
        t = OxmlElement("w:t")
        t.text = placeholder
        parts.append(run_with(t))
    parts.append(fld("end"))
    for r in parts:
        paragraph._p.append(r)


# Нумерация: каждый список — со своего номера. Отступы (twips): номер на 640,
# текст на 1134 (2 см) — заметный зазор, список глубже красной строки.
ABS_NUM_ID = 250
ABS_BULLET_ID = 251
_LVL_IND = '<w:pPr><w:ind w:left="1134" w:hanging="494"/></w:pPr>'
_LVL_IND_BULLET = '<w:pPr><w:ind w:left="851" w:hanging="227"/></w:pPr>'


class Numbering:
    """Нумерация одного документа: счётчики не глобальные, сборки не мешают друг другу."""

    def __init__(self, doc):
        self.root = doc.part.numbering_part.element
        self.counter = 500
        self.bullet_id = None
        self.lists = {}  # list_id разметки → numId

    def _ensure_abstract(self, abs_id, lvl_inner):
        for an in self.root.findall(qn("w:abstractNum")):
            if an.get(qn("w:abstractNumId")) == str(abs_id):
                return
        el = parse_xml(
            '<w:abstractNum %s w:abstractNumId="%d"><w:multiLevelType w:val="singleLevel"/>'
            '<w:lvl w:ilvl="0">%s</w:lvl></w:abstractNum>' % (nsdecls("w"), abs_id, lvl_inner))
        first_num = self.root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(el)
        else:
            self.root.append(el)

    def _add_num(self, abs_id, override=""):
        self.counter += 1
        self.root.append(parse_xml(
            '<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/>%s</w:num>'
            % (nsdecls("w"), self.counter, abs_id, override)))
        return self.counter

    def for_list(self, list_id, start):
        """numId нумерованного списка (новый — со start, повторный — тот же)."""
        if list_id not in self.lists:
            self._ensure_abstract(
                ABS_NUM_ID,
                '<w:start w:val="1"/><w:numFmt w:val="decimal"/>'
                '<w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>' + _LVL_IND)
            self.lists[list_id] = self._add_num(
                ABS_NUM_ID,
                '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="%d"/></w:lvlOverride>' % start)
        return self.lists[list_id]

    def bullet(self):
        if self.bullet_id is None:
            self._ensure_abstract(
                ABS_BULLET_ID,
                '<w:start w:val="1"/><w:numFmt w:val="bullet"/>'
                '<w:lvlText w:val="&#8226;"/><w:lvlJc w:val="left"/>' + _LVL_IND_BULLET)
            self.bullet_id = self._add_num(ABS_BULLET_ID)
        return self.bullet_id


def apply_numbering(paragraph, nid):
    paragraph._p.get_or_add_pPr().append(parse_xml(
        '<w:numPr %s><w:ilvl w:val="0"/><w:numId w:val="%d"/></w:numPr>'
        % (nsdecls("w"), nid)))


def add_plain(paragraph, text, size=None, bold=False):
    """Обычный текст; стрелки «→» в цепочках («Настройки → Конфиденциальность»)
    подсвечиваются акцентным цветом."""
    for piece in re.split(r"(→)", text):
        if not piece:
            continue
        run = paragraph.add_run(piece)
        run.bold = bold
        if piece == "→":
            run.bold = True
            run.font.color.rgb = ACCENT_RGB
        if size:
            run.font.size = Pt(size)


def add_inline(paragraph, text, size=None):
    """**b** *i* `code` [текст](url) → runs абзаца."""
    for kind, txt, url in mdparse.parse_inline(text):
        if kind == "link":
            if url.startswith(("http://", "https://")):
                add_hyperlink(paragraph, txt, url)
            else:  # относительная ссылка в docx не работает — просто текстом
                add_plain(paragraph, txt, size)
        elif kind == "bold":
            add_plain(paragraph, txt, size, bold=True)
        elif kind == "italic":
            run = paragraph.add_run(txt)
            run.italic = True
            if size:
                run.font.size = Pt(size)
        elif kind == "code":
            run = paragraph.add_run(txt)
            run.font.name = MONO
            run.font.size = Pt((size or 11) - 1.5)
            run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
            no_proof(run)
        else:
            add_plain(paragraph, txt, size)


def add_segments(paragraph, segs, size=None):
    """Сегменты абзаца: каждый следующий — с новой строки того же абзаца."""
    for k, seg in enumerate(segs):
        if k:
            paragraph.add_run().add_break()
        add_inline(paragraph, seg, size=size)
