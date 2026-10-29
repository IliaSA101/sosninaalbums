# -*- coding: utf-8 -*-
"""Тесты движка инструкции. Запуск: python -I -m unittest discover -s scripts/guide -v"""
import json
import re
import struct
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image
from PIL.PngImagePlugin import PngInfo

import annotate
import build_docx
import docx_xml
import mdparse
import place_run

REPO = Path(__file__).resolve().parents[2]
GUIDE = REPO / "docs" / "install" / "guide.md"
MAX_FILE = 400 * 1024
MAX_ASSETS = 6 * 1024 * 1024
MAX_VIDEO = 8 * 1024 * 1024  # ролик (решение 113)
MAX_VIDEOS = 30 * 1024 * 1024  # папка docs/install/video
# ссылка на ролик в guide.md — адрес raw/v{{VERSION}} (метка версии релиза), файл лежит в docs/install/video
VIDEO_RE = re.compile(r"/raw/v{{VERSION}}/docs/install/(video/[\w.-]+\.mp4)")
# Куски PNG, допустимые в снимках инструкции: ничего про автора, путь или время
SAFE_CHUNKS = {"IHDR", "PLTE", "tRNS", "IDAT", "IEND"}


def blocks(text: str):
    return mdparse.parse_blocks(text.splitlines())


def png_chunks(path: Path):
    data = path.read_bytes()
    pos, names = 8, []
    while pos < len(data):
        (size,) = struct.unpack(">I", data[pos:pos + 4])
        names.append(data[pos + 4:pos + 8].decode("ascii"))
        pos += 12 + size
    return names


def make_png(path: Path, size=(120, 80), color="white", meta=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", size, color)
    if meta:
        info = PngInfo()
        info.add_text("Author", "Иван Иванов C:/Users/ivan")
        img.save(path, "PNG", pnginfo=info, dpi=(300, 300))
    else:
        img.save(path, "PNG")


class TestMd(unittest.TestCase):
    def test_list_numbering_survives_image(self):
        bl = blocks("1. первый\n\n![рис](a.png)\n\n2. второй\n\nабзац\n\n1. новый список\n")
        ol = [b for b in bl if b.kind == "ol"]
        self.assertEqual([b.list_id for b in ol], [1, 1, 2])
        self.assertEqual([b.start for b in ol], [1, 2, 1])
        self.assertEqual([b.kind for b in bl].count("image"), 1)

    def test_tilde_lines(self):
        bl = blocks("абзац из\nдвух строк\n~ **Путь:** `x`\n~ **Ссылка:** y\n\n"
                    "1. пункт\n   с продолжением\n   ~ вторая строка\n")
        self.assertEqual(bl[0].segs, ("абзац из двух строк", "**Путь:** `x`", "**Ссылка:** y"))
        self.assertEqual(bl[1].segs, ("пункт с продолжением", "вторая строка"))

    def test_blocks_and_front_matter(self):
        meta, body = mdparse.parse_front_matter(
            ["---", "title: Т|Два", "toc: false", "---", "# Заг", "- а", "* б", "---", "\\page"])
        self.assertEqual(meta, {"title": "Т|Два", "toc": "false"})
        self.assertEqual([b.kind for b in mdparse.parse_blocks(body)],
                         ["heading", "ul", "ul", "hr", "page"])

    def test_table_and_code(self):
        bl = blocks("| а | б |\n|---|---|\n| 1 | 2 |\n\n```\nкод\n```\n")
        self.assertEqual(bl[0].header, ("а", "б"))
        self.assertEqual(bl[0].rows, (("1", "2"),))
        self.assertEqual(bl[1].kind, "code")

    def test_inline(self):
        got = mdparse.parse_inline("a **b** *c* `d` [e](https://x.y) z")
        self.assertEqual([k for k, _, _ in got],
                         ["text", "bold", "text", "italic", "text", "code", "text", "link", "text"])
        self.assertEqual(got[7], ("link", "e", "https://x.y"))

    def test_caption_prefix_stripped(self):
        self.assertEqual(mdparse.caption_text("Рис. 3. Окно"), "Окно")
        self.assertEqual(mdparse.caption_text("Окно"), "Окно")


class TestBuild(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def write_guide(self, body: str, pics=("a.png", "b.png")) -> Path:
        for p in pics:
            make_png(self.tmp / "assets" / p)
        md = self.tmp / "guide.md"
        md.write_text("---\ntitle: Тест\nversion: {{VERSION}}\ntoc: false\n---\n" + body,
                      encoding="utf-8")
        return md

    def test_no_unresolved_labels(self):
        md = self.write_guide("# Тест\n\nВерсия {{VERSION}}.\n\n![Рис. 7. Первое](assets/a.png)\n")
        out = build_docx.build(md, "1.2.3", self.tmp / "out")
        self.assertEqual(out.name, "SosninaAlbums-1.2.3-install-guide-ru.docx")
        xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
        self.assertNotIn("{{", xml)
        self.assertIn("Версия 1.2.3.", xml)
        # чужая метка — ошибка сборки, а не молчаливый текст в документе
        bad = self.write_guide("Метка {{PORTAL_ID}} осталась.\n")
        with self.assertRaises(build_docx.BuildError) as cm:
            build_docx.build(bad, "1.2.3", self.tmp / "out2")
        self.assertIn("{{PORTAL_ID}}", str(cm.exception))

    def test_figures_numbered_in_order(self):
        md = self.write_guide("![Рис. 9. Первое](assets/a.png)\n\n![второе](assets/b.png)\n")
        out = build_docx.build(md, "1.0.0", self.tmp / "out")
        caps = [p.text for p in Document(out).paragraphs if p.text.startswith("Рис.")]
        self.assertEqual(caps, ["Рис. 1. Первое", "Рис. 2. второе"])
        self.assertEqual(len(Document(out).inline_shapes), 2)

    def test_xml_child_order(self):
        """Word придирчив к порядку pPr/rPr/tblPr: после сборки он как в схеме."""
        md = self.write_guide("> врезка\n\n1. пункт\n   ~ строка `код`\n\n```\nx\n```\n\n"
                              "| а | б |\n|---|---|\n| `1` | 2 |\n\n![р](assets/a.png)\n")
        out = build_docx.build(md, "1.0.0", self.tmp / "out")
        root = Document(out).element
        for name, order in docx_xml._ORDER.items():
            idx = {n: i for i, n in enumerate(order.split())}
            for el in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + name):
                pos = [idx.get(k.tag.rsplit("}", 1)[-1], len(idx)) for k in el]
                self.assertEqual(pos, sorted(pos), (name, [k.tag for k in el]))

    def test_missing_image_fails(self):
        md = self.write_guide("![нет](assets/none.png)\n")
        with self.assertRaises(build_docx.BuildError) as cm:
            build_docx.build(md, "1.0.0", self.tmp / "out")
        self.assertIn("none.png", str(cm.exception))

    def test_cli_exit_codes(self):
        ok = self.write_guide("Текст.\n")
        self.assertEqual(build_docx.main([str(ok), "--version", "1.0.0",
                                          "--out", str(self.tmp / "o")]), 0)
        bad = self.write_guide("{{X}}\n")
        self.assertEqual(build_docx.main([str(bad), "--version", "1.0.0",
                                          "--out", str(self.tmp / "o2")]), 1)
        self.assertEqual(build_docx.main([str(ok), "--version", "latest",
                                          "--out", str(self.tmp / "o3")]), 1)

    def test_images_exist_and_budget(self):
        """Настоящий docs/install/guide.md: картинки на месте и в бюджете размера."""
        text = GUIDE.read_text(encoding="utf-8")
        refs = [m.group(2) for line in text.splitlines()
                if (m := mdparse.IMG_RE.match(line))]
        self.assertTrue(refs, "в guide.md нет ни одной картинки")
        for rel in refs:
            p = GUIDE.parent / rel
            self.assertTrue(p.is_file(), f"нет картинки {rel}")
            self.assertLessEqual(p.stat().st_size, MAX_FILE, f"{rel} больше 400 КБ")
        assets = GUIDE.parent / "assets"
        total = sum(f.stat().st_size for f in assets.rglob("*") if f.is_file())
        self.assertLessEqual(total, MAX_ASSETS, "папка assets больше 6 МБ")
        # ролики: каждая ссылка guide.md ведёт на файл в video/, файл и папка в бюджете
        videos = sorted(set(VIDEO_RE.findall(text)))
        self.assertTrue(videos, "в guide.md нет ссылок на ролики")
        for rel in videos:
            p = GUIDE.parent / rel
            self.assertTrue(p.is_file(), f"нет ролика {rel}")
            self.assertLessEqual(p.stat().st_size, MAX_VIDEO, f"{rel} больше 8 МБ")
        vdir = GUIDE.parent / "video"
        vtotal = sum(f.stat().st_size for f in vdir.rglob("*") if f.is_file())
        self.assertLessEqual(vtotal, MAX_VIDEOS, "папка video больше 30 МБ")

    def test_real_guide_builds(self):
        out = build_docx.build(GUIDE, "0.0.1", self.tmp / "out")
        xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
        self.assertNotIn("{{", xml)

    def test_variant_notes_only_in_mac_files(self):
        """«Этот файл — для macOS…» и повторные ссылки на ролики версии — только в файлах Mac, не в общем."""
        lines = GUIDE.read_text(encoding="utf-8").splitlines()
        common = "\n".join(mdparse.select_variant(lines))
        self.assertNotIn("Этот файл — для **macOS", common)
        for v in ("26", "15", "12-14"):
            text = "\n".join(mdparse.select_variant(lines, v))
            self.assertIn("Этот файл — для **macOS", text, v)

    def test_real_guide_every_screenshot_in_docx(self):
        """Каждый снимок guide.md попал в docx: рисунков столько же, у каждого подпись «Рис. N»."""
        text = GUIDE.read_text(encoding="utf-8")
        refs = [m.group(2) for line in text.splitlines() if (m := mdparse.IMG_RE.match(line))]
        out = build_docx.build(GUIDE, "0.0.1", self.tmp / "out")
        doc = Document(out)
        self.assertEqual(len(doc.inline_shapes), len(refs))
        caps = [p.text for p in doc.paragraphs if p.text.startswith("Рис.")]
        self.assertEqual(len(caps), len(refs))

    def test_figure_width_full_but_not_upscaled(self):
        full = int(build_docx.IMG_MAX_W)
        # полный экран — на всю ширину текста
        self.assertEqual(build_docx.figure_width(1920, 1080), full)
        self.assertEqual(build_docx.figure_width(1024, 768), full)
        # крупный план окна мельче ширины страницы при 96 ppi — не растягиваем (не размыть)
        small = build_docx.figure_width(384, 374)
        self.assertEqual(small, int(build_docx.Inches(1) * 384 / build_docx.IMG_MIN_PPI))
        self.assertLess(small, full)
        # очень высокий снимок ограничен по высоте
        tall = build_docx.figure_width(1000, 3000)
        self.assertLessEqual(tall * 3, int(build_docx.IMG_MAX_H) + 3)

    def test_docx_picture_size_matches_rule(self):
        make_png(self.tmp / "assets" / "big.png", size=(1920, 1080))
        make_png(self.tmp / "assets" / "win.png", size=(384, 374))
        md = self.tmp / "guide.md"
        md.write_text("---\ntitle: Т\ntoc: false\n---\n![a](assets/big.png)\n\n![b](assets/win.png)\n",
                      encoding="utf-8")
        out = build_docx.build(md, "1.0.0", self.tmp / "out")
        widths = [int(s.width) for s in Document(out).inline_shapes]
        self.assertEqual(widths[0], int(build_docx.IMG_MAX_W))
        self.assertEqual(widths[1], build_docx.figure_width(384, 374))


class TestAnnotate(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def run_one(self, spec: dict, **png_args) -> Path:
        raw = self.tmp / "raw" / "x.png"
        make_png(raw, **png_args)
        (self.tmp / "raw" / "x.ann.json").write_text(json.dumps(spec), encoding="utf-8")
        self.assertTrue(annotate.annotate_one(raw))
        return self.tmp / "annot" / "x.png"

    def test_mask_fills_rect(self):
        # рамка рисуется вне маски; стрелка не заходит в прямоугольник маски
        out = self.run_one({
            "masks": [{"rect": [10, 10, 49, 29], "color": "#808080"}],
            "boxes": [{"rect": [60, 40, 110, 70], "width": 2}],
        })
        with Image.open(out) as im:
            rgb = im.convert("RGB")
        for x, y in ((10, 10), (30, 20), (49, 29)):
            self.assertEqual(rgb.getpixel((x, y)), (128, 128, 128), (x, y))
        self.assertEqual(rgb.getpixel((5, 5)), (255, 255, 255))
        self.assertNotEqual(rgb.getpixel((60, 40)), (255, 255, 255))  # рамка нарисована

    def test_mask_under_marks(self):
        # маска раньше стрелок: красная линия поверх серой заливки остаётся видимой
        out = self.run_one({
            "masks": [{"rect": [0, 0, 119, 79], "color": "#808080"}],
            "arrows": [{"from": [10, 40], "to": [110, 40], "width": 4}],
        })
        with Image.open(out) as im:
            rgb = im.convert("RGB")
        r, g, b = rgb.getpixel((60, 40))
        self.assertTrue(r > 180 and g < 80 and b < 80, (r, g, b))  # красный, не серый
        self.assertEqual(rgb.getpixel((60, 10)), (128, 128, 128))

    def test_output_has_no_metadata(self):
        out = self.run_one({"boxes": [{"rect": [5, 5, 50, 50]}]}, meta=True)
        self.assertLessEqual(set(png_chunks(out)), SAFE_CHUNKS)
        with Image.open(out) as im:
            self.assertEqual(im.info.get("dpi"), None)
            self.assertFalse(im.getexif())
        # то же для всех готовых снимков инструкции в репозитории
        for p in (REPO / "docs" / "install" / "assets" / "annot").glob("*.png"):
            self.assertLessEqual(set(png_chunks(p)), SAFE_CHUNKS, p.name)

    def test_target_draws_box_label_and_arrow(self):
        # цель в левом верхнем углу большого снимка: подпись уходит вниз/вправо, не на цель
        out = self.run_one({"targets": [{"rect": [100, 100, 240, 130], "text": "Нажмите «Готово»"}]},
                           size=(1024, 768), color="#dddddd")
        with Image.open(out) as im:
            rgb = im.convert("RGB")
        pad = annotate.target_metrics(1024)["pad"]
        # рамка цели красная по границе (с отступом pad)
        r, g, b = rgb.getpixel((100 - pad, 115))
        self.assertTrue(r > 180 and g < 90 and b < 90, (r, g, b))
        # подпись и стрелка нарисованы: красные пиксели есть вне рамки
        red = [(x, y) for y in range(0, 768, 2) for x in range(0, 1024, 2)
               if (lambda p: p[0] > 180 and p[1] < 90 and p[2] < 90)(rgb.getpixel((x, y)))]
        self.assertGreater(len(red), 200)
        # внутри самой цели (без рамки) ничего не закрыто
        self.assertEqual(rgb.getpixel((170, 115)), (221, 221, 221))

    def test_place_label_inside_and_off_targets(self):
        img = Image.new("RGB", (400, 300), "white")
        for tgt in ([0, 0, 40, 20], [360, 280, 400, 300], [180, 130, 220, 160]):
            lab, side = annotate.place_label(img, tgt, (150, 40), 30, [tuple(tgt)])
            self.assertTrue(0 <= lab[0] and 0 <= lab[1] and lab[2] <= 400 and lab[3] <= 300, (tgt, lab))
            self.assertFalse(annotate._inter(lab, tgt), (tgt, lab, side))

    def test_place_label_prefers_flat_background(self):
        # сверху от цели пестрый фон, снизу ровный: подпись — вниз
        img = Image.new("RGB", (400, 300), "white")
        px = img.load()
        for y in range(0, 110):
            for x in range(400):
                px[x, y] = (0, 0, 0) if (x // 4 + y // 4) % 2 else (255, 255, 255)
        lab, side = annotate.place_label(img, [150, 120, 250, 150], (120, 30), 30, [(150, 120, 250, 150)])
        self.assertEqual(side, "below")

    def test_wrap_text_keeps_lines_within_width(self):
        from PIL import ImageFont
        font = ImageFont.truetype(annotate.find_font(), 20)
        wrapped = annotate.wrap_text("Нажмите кнопку «Все равно открыть» в разделе Безопасность\nвторая", font, 200)
        lines = wrapped.split("\n")
        self.assertGreater(len(lines), 2)
        self.assertTrue(all(font.getlength(x) <= 200 for x in lines))
        self.assertEqual(lines[-1], "вторая")

    def test_two_targets_do_not_overlap_labels(self):
        spec = {"targets": [
            {"rect": [300, 300, 340, 320], "text": "Первая подпись"},
            {"rect": [320, 330, 360, 350], "text": "Вторая подпись"},
        ]}
        out = self.run_one(spec, size=(800, 600), color="#cccccc")
        self.assertTrue(out.is_file())

    def test_skip_without_ann(self):
        raw = self.tmp / "raw" / "y.png"
        make_png(raw)
        self.assertFalse(annotate.annotate_one(raw))


class TestVariants(unittest.TestCase):
    """Файлы под версию macOS (`<!-- only: … -->`, build_docx --variant)."""

    TEXT = ["общее", "<!-- only: win -->", "виндовс", "<!-- /only -->",
            "<!-- only: mac -->", "для-одной-версии", "<!-- /only -->",
            "<!-- only: 15, 26 -->", "новые", "<!-- only: general -->", "в-общем-новые",
            "<!-- /only -->", "<!-- /only -->", "<!-- only: 12-14 -->", "старые", "<!-- /only -->"]

    def test_general_shows_all_but_mac_only(self):
        self.assertEqual(mdparse.select_variant(self.TEXT),
                         ["общее", "виндовс", "новые", "в-общем-новые", "старые"])

    def test_variant_filters_windows_other_versions_and_general(self):
        self.assertEqual(mdparse.select_variant(self.TEXT, "26"), ["общее", "для-одной-версии", "новые"])
        self.assertEqual(mdparse.select_variant(self.TEXT, "12-14"), ["общее", "для-одной-версии", "старые"])

    def test_win_variant_hides_mac_and_general(self):
        text = self.TEXT + ["<!-- only: general, mac -->", "раздел-мак", "<!-- /only -->",
                            "<!-- only: windows -->", "файл-виндовс", "<!-- /only -->"]
        self.assertEqual(mdparse.select_variant(text, "win"), ["общее", "виндовс", "файл-виндовс"])
        # раздел Mac (general, mac) виден в общем файле и в файле версии Mac; `windows` — только в файле Windows
        self.assertEqual(mdparse.select_variant(text)[-1], "раздел-мак")
        self.assertIn("раздел-мак", mdparse.select_variant(text, "15"))
        self.assertNotIn("файл-виндовс", mdparse.select_variant(text, "15"))

    def test_unbalanced_markers_fail(self):
        for bad in (["<!-- only: win -->", "x"], ["x", "<!-- /only -->"]):
            with self.assertRaises(ValueError):
                mdparse.select_variant(bad)

    def test_variant_image_prefers_version_suffix(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d) / "mac-04.png"
            make_png(base)
            self.assertEqual(build_docx.variant_image(base, "26"), base)  # снимка -26 нет — общий
            make_png(Path(d) / "mac-04-26.png")
            self.assertEqual(build_docx.variant_image(base, "26").name, "mac-04-26.png")
            self.assertEqual(build_docx.variant_image(base, "15"), base)  # у 15 суффикса нет
            win = Path(d) / "mac-04-win.png"
            make_png(win)
            self.assertEqual(build_docx.variant_image(win, "26"), win)
            make_png(Path(d) / "mac-04-26-win.png")  # крупный план: суффикс версии стоит перед «-win»
            self.assertEqual(build_docx.variant_image(win, "26").name, "mac-04-26-win.png")
            self.assertEqual(build_docx.variant_image(base, None), base)

    def test_names_and_unknown_variant(self):
        self.assertEqual(build_docx.out_name("1.2.3", "docx", "12-14"),
                         "SosninaAlbums-1.2.3-install-guide-macos-12-14-ru.docx")
        self.assertEqual(build_docx.out_name("1.2.3", "docx", "win"),
                         "SosninaAlbums-1.2.3-install-guide-windows-ru.docx")
        with tempfile.TemporaryDirectory() as d, self.assertRaises(build_docx.BuildError):
            build_docx.build(GUIDE, "1.0.0", Path(d), "13")

    def test_real_guide_builds_every_variant(self):
        with tempfile.TemporaryDirectory() as d:
            for v in (k for k in build_docx.VARIANTS if k not in ("win", "macos", "linux")):
                out = build_docx.build(GUIDE, "0.0.1", Path(d), v)
                self.assertIn(f"macos-{v}", out.name)
                doc = Document(out)
                text = "\n".join(p.text for p in doc.paragraphs)
                self.assertIn(build_docx.VARIANTS[v]["label"], text, v)
                self.assertNotIn("SmartScreen", text, v)  # Windows в файле про Mac не нужен
                self.assertNotIn("{{", text)

    def test_label_table(self):
        """Таблица меток design 039 §7.1: что видно в каком варианте."""
        text = []
        for tag in ("general", "both", "win", "windows", "linux", "mac", "26", "15", "12-14", "macos"):
            text += [f"<!-- only: {tag} -->", tag, "<!-- /only -->"]
        seen = lambda v: set(mdparse.select_variant(text, v))
        self.assertEqual(seen(None), {"general", "both", "win", "26", "15", "12-14"})
        self.assertEqual(seen("win"), {"win", "windows"})
        self.assertEqual(seen("macos"), {"general", "macos", "26", "15", "12-14"})
        self.assertEqual(seen("linux"), {"linux"})
        self.assertEqual(seen("26"), {"mac", "26"})
        self.assertEqual(seen("15"), {"mac", "15"})
        self.assertEqual(seen("12-14"), {"mac", "12-14"})

    def test_real_guide_general_has_no_own_file_lines(self):
        """Общий файл не показывает строки «для своего файла» (баг 0.10)."""
        text = "\n".join(mdparse.select_variant(GUIDE.read_text(encoding="utf-8").splitlines()))
        self.assertNotIn("Этот файл — для", text)

    def test_real_guide_macos_file(self):
        """Файл «все Mac»: три версии macOS есть, Windows нет."""
        self.assertEqual(build_docx.out_name("1.2.3", "docx", "macos"),
                         "SosninaAlbums-1.2.3-install-guide-macos-ru.docx")
        with tempfile.TemporaryDirectory() as d:
            out = build_docx.build(GUIDE, "0.0.1", Path(d), "macos")
            doc = Document(out)
            text = "\n".join(p.text for p in doc.paragraphs)
            tables = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn("Как установить на Mac", text)
            self.assertIn("macOS 12–26", text)
            for v in ("macOS 15 (Sequoia) и 26 (Tahoe)", "macOS 12, 13 и 14"):
                self.assertIn(v, text)
            self.assertIn("Tahoe", text)
            self.assertIn("Monterey", text)
            for win in ("SmartScreen", "Windows 10", ".exe", "Этот файл — для"):
                self.assertNotIn(win, text)
                self.assertNotIn(win, tables)
            self.assertIn("install-guide-macos-ru.docx", tables)
            self.assertNotIn("{{", text)

    def test_real_guide_windows_file(self):
        """Файл «только Windows»: SmartScreen и рисунки Windows есть, Mac нет."""
        with tempfile.TemporaryDirectory() as d:
            out = build_docx.build(GUIDE, "0.0.1", Path(d), "win")
            self.assertEqual(out.name, "SosninaAlbums-0.0.1-install-guide-windows-ru.docx")
            doc = Document(out)
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("Windows 10 и 11", text)
            self.assertIn("SmartScreen", text)
            for mac in ("Gatekeeper", "Finder", "xattr", "macOS", "Терминал"):
                self.assertNotIn(mac, text)
            refs = [m.group(2) for line in mdparse.select_variant(GUIDE.read_text(encoding="utf-8").splitlines(), "win")
                    if (m := mdparse.IMG_RE.match(line))]
            self.assertTrue(refs and all("/win-" in r for r in refs), refs)
            self.assertEqual(len(doc.inline_shapes), len(refs))


    def test_real_guide_linux_file(self):
        """101-13: файл Linux собирается; в нём нет разделов Windows и Mac, в общих файлах нет Linux."""
        self.assertEqual(build_docx.out_name("1.2.3", "docx", "linux"),
                         "SosninaAlbums-1.2.3-install-guide-linux-ru.docx")
        with tempfile.TemporaryDirectory() as d:
            out = build_docx.build(GUIDE, "0.0.1", Path(d), "linux")
            self.assertEqual(out.name, "SosninaAlbums-0.0.1-install-guide-linux-ru.docx")
            doc = Document(out)
            text = "\n".join(p.text for p in doc.paragraphs)
            tables = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn("Как установить на Linux", text)
            for want in ("sudo apt install ./SosninaAlbums-<версия>-linux-amd64.deb", "sudo apt remove sosninaalbums",
                         "GTK 3", "WebKitGTK 4.1"):
                self.assertIn(want, text)
            for want in ("~/.config/SosninaAlbums", "~/.cache/SosninaAlbums", "~/.local/share/SosninaAlbums/backups"):
                self.assertIn(want, tables)
            self.assertIn("linux-amd64.tar.gz", tables)
            for other in ("SmartScreen", "Windows 10", ".exe", "Gatekeeper", "Finder", "xattr", "macOS", "{{"):
                self.assertNotIn(other, text)
                self.assertNotIn(other, tables)
            self.assertEqual(len(doc.inline_shapes), 0)  # снимков экрана Linux в 0.23 нет
        # раздел Linux не попадает ни в общий файл, ни в файлы Windows и Mac
        lines = GUIDE.read_text(encoding="utf-8").splitlines()
        for v in (None, "win", "macos", "26", "15", "12-14"):
            self.assertNotIn("linux-amd64", "\n".join(mdparse.select_variant(lines, v)), v)
            self.assertNotIn("apt install", "\n".join(mdparse.select_variant(lines, v)), v)


class TestPlaceRun(unittest.TestCase):
    def test_macos14_run_gets_suffix_and_steps_10_11(self):
        with tempfile.TemporaryDirectory() as d:
            run, assets = Path(d) / "run", Path(d) / "assets"
            make_png(run / "png" / "10-context-menu.png", size=(800, 500))
            make_png(run / "png" / "11-open-dialog.png", size=(800, 500))
            make_png(run / "png" / "11-open-dialog-win.png", size=(300, 200))
            shots = [
                {"name": "10-context-menu", "ok": True, "file": "png/10-context-menu.png",
                 "target": {"rect": [100, 100, 60, 60], "label": "sosninaalbums"}},
                {"name": "11-open-dialog", "ok": True, "file": "png/11-open-dialog.png",
                 "win": "png/11-open-dialog-win.png", "window": [250, 150, 300, 200],
                 "target": {"rect": [450, 300, 80, 24], "label": "Открыть"}},
            ]
            (run / "shots.json").write_text(json.dumps(shots), encoding="utf-8")
            place_run.place(run, "14", assets, {}, "тест")
            for name in ("mac-10-context-menu-14", "mac-11-open-dialog-14", "mac-11-open-dialog-14-win"):
                self.assertTrue((assets / "raw" / f"{name}.png").is_file(), name)
                self.assertTrue((assets / "annot" / f"{name}.png").is_file(), name)
            spec = json.loads((assets / "raw" / "mac-11-open-dialog-14.ann.json").read_text(encoding="utf-8"))
            self.assertEqual(spec["targets"][0]["text"], "Нажмите «Открыть»")

    def test_side_text_override_and_window_only(self):
        with tempfile.TemporaryDirectory() as d:
            run, assets = Path(d) / "run", Path(d) / "assets"
            make_png(run / "png" / "06-confirm.png", size=(800, 500))
            make_png(run / "png" / "07-terminal-xattr.png", size=(800, 500))
            make_png(run / "png" / "07-terminal-xattr-win.png", size=(300, 200))
            shots = [
                {"name": "06-confirm", "ok": True, "file": "png/06-confirm.png",
                 "target": {"rect": [450, 300, 80, 24], "label": "Все равно открыть"}},
                {"name": "07-terminal-xattr", "ok": True, "file": "png/07-terminal-xattr.png",
                 "win": "png/07-terminal-xattr-win.png", "window": [10, 10, 300, 200],
                 "target": {"rect": [20, 20, 30, 10], "label": "x"}},
            ]
            (run / "shots.json").write_text(json.dumps(shots), encoding="utf-8")
            over = {"06-confirm": {"text": "Жмите «{label}»", "side": "left"}}
            place_run.place(run, "15", assets, over, "тест")
            spec = json.loads((assets / "raw" / "mac-06-confirm.ann.json").read_text(encoding="utf-8"))
            self.assertEqual(spec["targets"][0]["side"], "left")  # override важнее STEPS
            self.assertEqual(spec["targets"][0]["text"], "Жмите «Все равно открыть»")
            # окно «Терминала» вместо экрана: размер окна, без стрелки
            from PIL import Image
            with Image.open(assets / "raw" / "mac-07-terminal-xattr.png") as im:
                self.assertEqual(im.size, (300, 200))
            spec7 = json.loads((assets / "raw" / "mac-07-terminal-xattr.ann.json").read_text(encoding="utf-8"))
            self.assertEqual(spec7, {})


if __name__ == "__main__":
    unittest.main()
